import asyncio
import io
import logging
import os
import shutil
import subprocess
import tempfile
from typing import Any, Dict, List

import docx
import fitz
import ocrmypdf
import requests
from fastapi import HTTPException

from app.services.llm_validator import extract_items_from_text

logger = logging.getLogger(__name__)

BASE_URL = 'https://public.api.openprocurement.org/api/2.5'


def get_contract(contract_id):
    url = f'{BASE_URL}/tenders/{contract_id}'

    try:
        response = requests.get(url)
        response.raise_for_status()
        return response.json().get("data", {})
    except requests.HTTPError as err:
        raise HTTPException(status_code=response.status_code, detail=f"Error: {err}")


_SUPPORTED_FORMATS = {
    'application/pdf': 'pdf',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document': 'docx',
    'application/msword': 'doc',
}


def get_document_urls(documents):
    """Return list of (url, doc_type) tuples for supported public documents."""
    result = []
    for doc in documents:
        if doc.get('confidentiality') != 'public':
            continue
        fmt = doc.get('format', '')
        doc_type = _SUPPORTED_FORMATS.get(fmt)
        if doc_type:
            result.append((doc.get('url'), doc_type))
    if not result:
        raise Exception("There are no supported documents attached to tender!")
    return result


def ocr_pdf(input_file: str, output_file: str, lang: str = "ukr+eng"):
    """Run OCR on a single PDF file, writing the result to output_file."""
    ocrmypdf.ocr(
        input_file=input_file,
        output_file=output_file,
        language=lang,
        tesseract_pagesegmode=3,
        tesseract_oem=1,
        force_ocr=True,
        output_type="pdf",
        deskew=True,
        rotate_pages=True,
        pdf_renderer="hocr",
    )


def _is_scanned_page(page) -> bool:
    """Determine if a PDF page is a scanned image (needs OCR) vs text-based.

    Uses fitz block analysis: if text blocks cover >5% of the page area,
    the page has a usable text layer and OCR can be skipped.
    """
    page_area = page.rect.width * page.rect.height
    if page_area <= 0:
        return True

    blocks = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]

    text_area = 0
    image_area = 0

    for block in blocks:
        bx0 = block.get("bbox", (0, 0, 0, 0))
        block_area = abs((bx0[2] - bx0[0]) * (bx0[3] - bx0[1]))

        if block.get("type") == 0:  # text block
            # Check if the block actually has meaningful text
            has_text = False
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if span.get("text", "").strip():
                        has_text = True
                        break
                if has_text:
                    break
            if has_text:
                text_area += block_area
        elif block.get("type") == 1:  # image block
            image_area += block_area

    text_coverage = text_area / page_area

    # If text blocks cover more than 5% of the page, it's a text PDF
    if text_coverage > 0.05:
        return False

    # If the page is mostly images, it's likely a scan
    return True


def _extract_text_from_page(page_bytes: bytes, page_num: int, work_dir: str) -> str:
    """Extract text from a single PDF page. Uses smart detection to skip OCR for text PDFs."""
    doc = fitz.open(stream=page_bytes, filetype="pdf")
    page = doc[0]

    # Smart detection: check if page has enough text blocks
    if not _is_scanned_page(page):
        text = page.get_text()
        doc.close()
        if len(text.strip()) > 50:
            logger.info(f"Page {page_num}: text-based PDF, OCR skipped ({len(text)} chars)")
            return text

    # Fallback: try basic text extraction
    text = page.get_text()
    doc.close()

    if len(text.strip()) > 100:
        return text

    # Text layer is sparse — run OCR
    logger.info(f"Page {page_num}: scanned page detected, running OCR")
    page_path = os.path.join(work_dir, f"page_{page_num}.pdf")
    ocr_path = os.path.join(work_dir, f"page_{page_num}_ocred.pdf")

    with open(page_path, "wb") as f:
        f.write(page_bytes)

    try:
        ocr_pdf(page_path, ocr_path)
        ocr_doc = fitz.open(ocr_path)
        ocr_text = ocr_doc[0].get_text()
        ocr_doc.close()
        return ocr_text if ocr_text.strip() else text
    except Exception as e:
        logger.warning(f"OCR failed for page {page_num}: {e}")
        return text


async def _extract_text_all_pages(pdf_content: bytes) -> str:
    """Extract text from all pages of a PDF in parallel (OCR only when needed)."""
    doc = fitz.open(stream=pdf_content, filetype="pdf")

    # Split into per-page byte buffers in memory
    page_buffers = []
    for i in range(len(doc)):
        single = fitz.open()
        single.insert_pdf(doc, from_page=i, to_page=i)
        page_buffers.append(single.tobytes())
        single.close()
    doc.close()

    work_dir = tempfile.mkdtemp(prefix="pyman_pdf_")
    try:
        tasks = [
            asyncio.to_thread(_extract_text_from_page, buf, i, work_dir)
            for i, buf in enumerate(page_buffers)
        ]
        results = await asyncio.gather(*tasks)
        return "\n".join(results)
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


async def parse_documents(document_refs: List[tuple]) -> List[Dict[str, Any]]:
    """Download PDFs/DOCX/DOC, extract text, and use LLM to extract items.

    document_refs: list of (url, doc_type) where doc_type is 'pdf', 'docx', or 'doc'.
    """
    all_text = []
    for url, doc_type in document_refs:
        try:
            response = requests.get(url)
            response.raise_for_status()
            if doc_type == 'pdf':
                text = await _extract_text_all_pages(response.content)
            elif doc_type == 'docx':
                text = extract_text_from_docx(response.content)
            elif doc_type == 'doc':
                pdf_bytes = await asyncio.to_thread(convert_doc_to_pdf, response.content)
                text = await _extract_text_all_pages(pdf_bytes)
            else:
                continue
            if text.strip():
                all_text.append(text)
        except Exception as e:
            logger.error(f"Failed to process {doc_type} from {url}: {e}")

    if not all_text:
        raise Exception("Could not extract text from any document!")

    combined_text = "\n\n".join(all_text)
    items = await extract_items_from_text(combined_text)

    if items is None:
        raise Exception("LLM failed to extract items from document text")

    return items


def convert_doc_to_pdf(doc_content: bytes) -> bytes:
    """Convert a legacy .doc file to PDF bytes via LibreOffice headless.

    Raises an exception if soffice is unavailable or conversion fails.
    """
    work_dir = tempfile.mkdtemp(prefix="pyman_doc_")
    try:
        src = os.path.join(work_dir, "input.doc")
        with open(src, "wb") as f:
            f.write(doc_content)

        result = subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to", "pdf",
                "--outdir", work_dir,
                src,
            ],
            capture_output=True,
            timeout=120,
        )
        if result.returncode != 0:
            raise Exception(
                f"LibreOffice conversion failed (code {result.returncode}): "
                f"{result.stderr.decode('utf-8', errors='ignore')[:300]}"
            )

        pdf_path = os.path.join(work_dir, "input.pdf")
        if not os.path.exists(pdf_path):
            raise Exception("LibreOffice did not produce a PDF output")

        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def validate_name(name, max_elements=2):
    items = name.split(",")

    if len(items) > max_elements:
        return False

    return True


async def get_contract_info(contract_id):
    contract = get_contract(contract_id)
    logger.info(f"Contract data for {contract_id}: {str(contract)[:200]}")
    if not contract:
        raise HTTPException(status_code=404, detail="Contract not found!")

    if 'items' in contract:
        items = []
        for item in contract['items']:
            name = item.get('description', ' ')
            quantity = item.get('quantity', ' ')
            unit_name = item.get('unit', {}).get('name', 'од.')
            unit_price = item.get('unit', {}).get('value', {}).get('amount', ' ')
            data = {'name': name, 'quantity': quantity, 'unit_name': unit_name, 'unit_price': unit_price}

            # if the last one price for unit is not defined, then it tries to calculate it automatically, if there is only one item,
            # otherwise it will parse the documents to find all information there
            if unit_price == ' ':
                # here i am validating name in response for better filtration, because if there are a lot of commas this name is incorrect, it might have a lot of element names in it,
                # for example without this filter response can be: [{'name': 'Господарські товари( грунт,стрічка малярна,пінопласт,ніж канцелярський, клей Перфлікс)',
                # 'quantity': 15.0, 'unit_name': 'штука', 'unit_price': 63.733333333333334, 'total_price': 956.0}], its bad response for selenium price parsing,
                # because it has not clear name
                if len(contract['items']) == 1 and validate_name(name):
                    total_price = contract.get("value", {}).get("amount", " ")
                    if not total_price == " ":
                        data['total_price'] = total_price
                        data["unit_price"] = float(total_price) / float(quantity)
                    items.append(data)
                    return items
                else:
                    if validate_name(name):
                        items.append(data)
                        continue
                    else:
                        break

            items.append(data)

    if len(items) > 0:
        return items

    # if price is not placed for the first element or if it cannot be calculated then program will parse the pdf to search the needed infomation
    documents = contract.get('documents', None)
    if documents:
        try:
            url = get_document_urls(documents)
            items = await parse_documents(url)
            if len(items) == 0:
                raise HTTPException(status_code=404, detail="Could not find items in specification!")
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=404, detail=str(e))
    else:
        raise HTTPException(status_code=404, detail="Information about tender could not be found!")

    return items


def get_tender_documents(contract_id):
    url = f"{BASE_URL}/tenders/{contract_id}/documents"
    try:
        response = requests.get(url)
        response.raise_for_status()
        docs = response.json().get("data", [])
        valid_docs = []
        for doc in docs:
            if doc.get('confidentiality') == 'public':
                title = doc.get('title', '').lower()
                if (title.endswith('.pdf') or title.endswith('.docx')) and not title.endswith('.p7s'):
                    valid_docs.append(doc)

        unique_docs = {}
        for d in valid_docs:
            unique_docs[d.get('title')] = d
        return list(unique_docs.values())
    except Exception as e:
        logger.warning(f"Error getting documents: {e}")
        return []


def extract_text_from_docx(content: bytes) -> str:
    doc = docx.Document(io.BytesIO(content))
    text = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text.strip())
    return "\n".join(text)


def extract_text_from_pdf_extended(content: bytes) -> str:
    """Extract text from PDF with selective OCR per page. Sync — called via asyncio.to_thread()."""
    doc = fitz.open(stream=content, filetype="pdf")
    work_dir = tempfile.mkdtemp(prefix="pyman_pdf_ext_")

    try:
        page_texts = []
        for i in range(len(doc)):
            # Get per-page bytes
            single = fitz.open()
            single.insert_pdf(doc, from_page=i, to_page=i)
            page_bytes = single.tobytes()
            single.close()

            page_texts.append(_extract_text_from_page(page_bytes, i, work_dir))

        return "\n".join(page_texts)
    finally:
        doc.close()
        shutil.rmtree(work_dir, ignore_errors=True)
